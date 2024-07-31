import streamlit as st
import os
import base64
from gtts import gTTS
import PyPDF2
import docx
import markdown
from PIL import Image
import pytesseract
from deep_translator import GoogleTranslator
from datetime import datetime
from fpdf import FPDF
from dotenv import load_dotenv
import groq
import ast

load_dotenv()
API_KEY = os.getenv("GROQ_API_KEY")

# Constants
PERSONAS = {
    "Default": "A default persona",
    # Add more personas as needed
}

# Initialize session state
def initialize_session_state():
    default_values = {
        "messages": [],
        "audio_base64": "",
        "file_content": "",
        "chat_histories": [],
        "persona": "Default",
        "custom_persona": "",
        "user": None,
        "model_params": {
            "model": "llama-3.1-70b-versatile",
            "max_tokens": 1024,
            "temperature": 0.7,
            "top_p": 0.9,
        },
        "total_tokens": 0,
        "total_cost": 0,
        "enable_audio": False,
        "language": "English",
        "show_analysis": False,
        "content_creation_mode": False,
        "show_summarization": False,
        "summarization_type": "Main Takeaways",
    }
    for key, value in default_values.items():
        if key not in st.session_state:
            st.session_state[key] = value

initialize_session_state()

# Groq client setup
def get_groq_client(api_key: str) -> groq.Groq:
    return groq.Groq(api_key=api_key)

# Stream LLM response
def stream_llm_response(
    client: groq.Groq, params: dict, messages: list
) -> str:
    try:
        chat_completion = client.chat.completions.create(
            messages=messages,
            model=params["model"],
            max_tokens=params["max_tokens"],
            temperature=params["temperature"],
            top_p=params["top_p"],
            stream=True,
        )
        for chunk in chat_completion:
            if chunk.choices[0].delta.content is not None:
                yield chunk.choices[0].delta.content
    except Exception as e:
        st.error(f"Error in API call: {str(e)}")
        yield "I'm sorry, but I encountered an error while processing your request."

# Validate prompt
def validate_prompt(prompt: str) -> None:
    if not prompt.strip():
        raise ValueError("Prompt cannot be empty")

# Analyze code
def analyze_code(code: str) -> list:
    try:
        tree = ast.parse(code)
        analysis_results = []
        for node in ast.walk(tree):
            if isinstance(node, ast.FunctionDef):
                function_name = node.name
                docstring = ast.get_docstring(node)
                arguments = [arg.arg for arg in node.args.args]
                analysis_results.append(
                    {
                        "Function": function_name,
                        "Docstring": docstring,
                        "Arguments": arguments,
                    }
                )
        return analysis_results
    except SyntaxError as e:
        return [{"Error": f"Syntax Error: {e.msg}"}]

# Summarize text
def summarize_text(text: str, summarization_type: str = "Main Takeaways") -> str:
    client = get_groq_client(API_KEY)
    messages = [
        {"role": "assistant", "content": f"Summarize the text in {summarization_type} format."},
        {"role": "user", "content": text},
    ]
    response = client.chat.completions.create(
        messages=messages,
        model=st.session_state.model_params["model"],
        max_tokens=st.session_state.model_params["max_tokens"],
    )
    return response.choices[0].message.content

# Create content
def create_content(prompt: str, content_type: str = "Story") -> str:
    client = get_groq_client(API_KEY)
    messages = [
        {"role": "assistant", "content": f"Create a {content_type} based on the prompt."},
        {"role": "user", "content": prompt},
    ]
    response = client.chat.completions.create(
        messages=messages,
        model=st.session_state.model_params["model"],
        max_tokens=st.session_state.model_params["max_tokens"],
    )
    return response.choices[0].message.content

# Export chat
def export_chat(format: str) -> None:
    if format == "PDF":
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=15)
        for message in st.session_state.messages:
            pdf.cell(200, 10, txt=message["role"] + ": " + message["content"], ln=True, align='L')
        pdf.output("chat.pdf")
        with open("chat.pdf", "rb") as file:
            st.download_button("Download PDF", file.read(), "chat.pdf")
    elif format == "Text":
        text = ""
        for message in st.session_state.messages:
            text += message["role"] + ": " + message["content"] + "\n"
        st.download_button("Download Text", text, "chat.txt")

# Process uploaded file
def process_uploaded_file(uploaded_file) -> str:
    if uploaded_file.type == "application/pdf":
        pdf_reader = PyPDF2.PdfReader(uploaded_file)
        text = ""
        for page in range(len(pdf_reader.pages)):
            text += pdf_reader.pages[page].extract_text()
        return text
    elif uploaded_file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        doc = docx.Document(uploaded_file)
        text = ""
        for para in doc.paragraphs:
            text += para.text + "\n"
        return text
    elif uploaded_file.type == "image/jpeg" or uploaded_file.type == "image/png":
        img = Image.open(uploaded_file)
        text = pytesseract.image_to_string(img)
        return text
    else:
        return "Unsupported file type"

# Save chat history
def save_chat_history() -> None:
    chat_history = {
        "messages": st.session_state.messages,
        "model_params": st.session_state.model_params,
        "total_tokens": st.session_state.total_tokens,
        "total_cost": st.session_state.total_cost,
    }
    st.session_state.chat_histories.append(chat_history)

# Load chat history
def load_chat_history(selected_history: str) -> None:
    for history in st.session_state.chat_histories:
        if history["messages"][0]["content"] == selected_history:
            st.session_state.messages = history["messages"]
            st.session_state.model_params = history["model_params"]
            st.session_state.total_tokens = history["total_tokens"]
            st.session_state.total_cost = history["total_cost"]
            break

# Update token count
def update_token_count(tokens: int) -> None:
    st.session_state.total_tokens += tokens

# Text to speech
def text_to_speech(text: str, lang: str) -> None:
    tts = gTTS(text=text, lang=lang)
    tts.save("speech.mp3")
    with open("speech.mp3", "rb") as file:
        st.download_button("Download Speech", file.read(), "speech.mp3")

# Reset all
def reset_all() -> None:
    st.session_state.messages = []
    st.session_state.audio_base64 = ""
    st.session_state.file_content = ""
    st.session_state.model_params = {
        "model": "llama-3.1-70b-versatile",
        "max_tokens": 1024,
        "temperature": 0.7,
        "top_p": 0.9,
    }
    st.session_state.total_tokens = 0
    st.session_state.total_cost = 0
    st.session_state.enable_audio = False
    st.session_state.language = "English"
    st.session_state.show_analysis = False
    st.session_state.content_creation_mode = False
    st.session_state.show_summarization = False
    st.session_state.summarization_type = "Main Takeaways"

# Translate text
def translate_text(text: str, target_lang: str) -> str:
    translator = GoogleTranslator(source="auto", target=target_lang)
    return translator.translate(text)

# Main Streamlit app
def main() -> None:
    st.title("Chatbot")
    st.session_state.user = st.selectbox("Select User", ["User", "Assistant"])

    if st.session_state.user == "User":
        prompt = st.text_area("Enter your prompt")
        if st.button("Submit"):
            validate_prompt(prompt)
            client = get_groq_client(API_KEY)
            messages = [
                {"role": "user", "content": prompt},
            ]
            response = client.chat.completions.create(
                messages=messages,
                model=st.session_state.model_params["model"],
                max_tokens=st.session_state.model_params["max_tokens"],
            )
            st.session_state.messages.append({"role": "user", "content": prompt})
            st.session_state.messages.append({"role": "assistant", "content": response.choices[0].message.content})
            update_token_count(response.usage.total_tokens)
    else:
        if st.button("Create Content"):
            st.session_state.content_creation_mode = True
        if st.session_state.content_creation_mode:
            prompt = st.text_area("Enter your prompt")
            content_type = st.selectbox("Select content type", ["Story", "Poem", "Dialogue"])
            if st.button("Generate Content"):
                content = create_content(prompt, content_type)
                st.write(content)
        if st.button("Summarize Text"):
            st.session_state.show_summarization = True
        if st.session_state.show_summarization:
            text = st.text_area("Enter your text")
            summarization_type = st.selectbox("Select summarization type", ["Main Takeaways", "Summary", "Outline"])
            if st.button("Summarize"):
                summary = summarize_text(text, summarization_type)
                st.write(summary)

    if st.button("Export Chat"):
        export_format = st.selectbox("Select export format", ["PDF", "Text"])
        export_chat(export_format)

    if st.button("Save Chat History"):
        save_chat_history()

    if st.button("Load Chat History"):
        selected_history = st.selectbox("Select chat history", [message["content"] for message in st.session_state.messages])
        load_chat_history(selected_history)

    if st.button("Reset All"):
        reset_all()

if __name__ == "__main__":
    main()